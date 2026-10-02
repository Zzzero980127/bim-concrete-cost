"""
BIM 混凝土成本自動估算：把 Forma Takeoff 匯出的體積 × 最新單價，產生網頁用的 data.json。

輸入（都在 data/ 資料夾）：
  quantities.xlsx 或 quantities.csv  ← 從 Forma Takeoff 匯出，模型改版時覆蓋這一份就好
  model_info.json                    ← 模型名稱與備註（選填）
  unit_prices.csv                    ← 單價來源（之後可改成爬蟲，見 fetch_market_prices）

自動維護（不要手動改）：
  versions.json       ← 偵測到 quantities 內容變動時，自動新增一個版本
  price_history.csv   ← 每次執行記錄當天單價，同一天只記一次

輸出：
  data.json           ← 給 index.html 讀取
"""

import csv
import hashlib
import json
import os
import re
import sys
from datetime import datetime, timedelta, timezone

sys.stdout.reconfigure(encoding="utf-8")  # Windows 終端機預設 cp950，印不出 m³

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(ROOT, "data")
OUTPUT_FILE = os.path.join(ROOT, "data.json")

TW = timezone(timedelta(hours=8))

# Takeoff 匯出的欄位名稱可能是中文或英文，這裡列出可接受的名稱（拿到實際匯出檔後再調整）
COLUMN_ALIASES = {
    "category": ["類別", "構件類別", "Category", "Classification", "Type"],
    "spec": ["材質", "規格", "材料", "Material", "Description"],
    "volume": ["體積", "體積(m3)", "體積 (m³)", "Volume", "Quantity"],
    "count": ["數量", "個數", "Count"],
}

# 把 Revit/Takeoff 的類別名稱歸到四大類
CATEGORY_KEYWORDS = {
    "floors": ["樓板", "板", "floor", "slab"],
    "columns": ["柱", "column"],
    "beams": ["梁", "樑", "framing", "beam"],
    "walls": ["牆", "wall"],
}

logs = []


def log(tag, msg):
    now = datetime.now(TW).strftime("%H:%M:%S")
    line = f"[{now}] {tag} {msg}"
    logs.append({"time": now, "tag": tag, "msg": msg})
    print(line)


def today():
    # BUILD_DATE 只用來測試（模擬不同日期執行），平常不用設定
    return os.environ.get("BUILD_DATE") or datetime.now(TW).strftime("%Y-%m-%d")


# ---------- 讀取 Takeoff 匯出檔 ----------

def read_table(path):
    if path.endswith(".xlsx"):
        from openpyxl import load_workbook
        ws = load_workbook(path, data_only=True).active
        rows = [[("" if c is None else str(c)).strip() for c in r] for r in ws.iter_rows(values_only=True)]
    else:
        with open(path, encoding="utf-8-sig", newline="") as f:
            rows = [[c.strip() for c in r] for r in csv.reader(f)]
    rows = [r for r in rows if any(r)]
    header, body = rows[0], rows[1:]
    return [dict(zip(header, r)) for r in body]


def pick(row, field):
    for name in COLUMN_ALIASES[field]:
        if name in row and row[name] != "":
            return row[name]
    return None


def classify(category):
    text = (category or "").lower()
    for key, words in CATEGORY_KEYWORDS.items():
        if any(w in text for w in words):
            return key
    return None


def parse_grade(spec):
    m = re.search(r"(\d{3})\s*kgf", spec or "")
    return m.group(1) if m else "280"


def to_number(value):
    return float(re.sub(r"[^\d.\-]", "", value or "0") or 0)


def load_quantities():
    for name in ("quantities.xlsx", "quantities.csv"):
        path = os.path.join(DATA_DIR, name)
        if os.path.exists(path):
            break
    else:
        raise FileNotFoundError("找不到 data/quantities.xlsx 或 data/quantities.csv")

    with open(path, "rb") as f:
        digest = hashlib.sha256(f.read()).hexdigest()[:12]

    elements = {}
    for row in read_table(path):
        key = classify(pick(row, "category"))
        if not key:
            continue
        spec = pick(row, "spec") or ""
        el = elements.setdefault(key, {"volume": 0.0, "count": 0, "spec": spec, "grade": parse_grade(spec)})
        el["volume"] += to_number(pick(row, "volume"))
        el["count"] += int(to_number(pick(row, "count") or "1"))

    for el in elements.values():
        el["volume"] = round(el["volume"], 2)

    log("TAKEOFF", f"讀取 {name}：{len(elements)} 類構件，總體積 {sum(e['volume'] for e in elements.values()):,.2f} m³")
    return name, digest, elements


# ---------- 單價 ----------

def fetch_market_prices():
    """
    取得最新單價，回傳 {"280": 3250, ...} 與來源說明。
    目前讀 data/unit_prices.csv（手動維護）。決定好外部價格網站後，
    把爬蟲寫在這裡、回傳相同格式即可，其他程式不用改。
    """
    prices, source = {}, ""
    with open(os.path.join(DATA_DIR, "unit_prices.csv"), encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            prices[row["grade"].strip()] = float(row["price"])
            source = row.get("source", "").strip() or source
    log("PRICE", "取得單價 " + "、".join(f"{g} kgf = NT$ {p:,.0f}" for g, p in sorted(prices.items())))
    return prices, source


def update_price_history(date, prices, source):
    path = os.path.join(DATA_DIR, "price_history.csv")
    rows = []
    if os.path.exists(path):
        with open(path, encoding="utf-8-sig", newline="") as f:
            rows = list(csv.DictReader(f))

    rows = [r for r in rows if r["date"] != date]  # 同一天重跑就覆蓋當天紀錄
    for grade, price in sorted(prices.items()):
        rows.append({"date": date, "grade": grade, "price": f"{price:g}", "source": source})
    rows.sort(key=lambda r: (r["date"], r["grade"]))

    with open(path, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["date", "grade", "price", "source"])
        writer.writeheader()
        writer.writerows(rows)

    history = {}
    for r in rows:
        history.setdefault(r["date"], {})[r["grade"]] = float(r["price"])
    log("PRICE", f"價格歷史共 {len(history)} 天")
    return history


# ---------- 版本 ----------

def update_versions(date, filename, digest, elements):
    path = os.path.join(DATA_DIR, "versions.json")
    versions = []
    if os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            versions = json.load(f)

    if versions and versions[-1]["hash"] == digest:
        log("MODEL", f"體積檔案沒有變動，沿用 {versions[-1]['version']}")
        return versions

    info = {}
    info_path = os.path.join(DATA_DIR, "model_info.json")
    if os.path.exists(info_path):
        with open(info_path, encoding="utf-8") as f:
            info = json.load(f)

    version = {
        "version": f"V{len(versions) + 1}",
        "model_name": info.get("model_name", filename),
        "note": info.get("note", ""),
        "author": info.get("author", ""),
        "detected_at": date,
        "hash": digest,
        "elements": elements,
    }
    versions.append(version)
    with open(path, "w", encoding="utf-8") as f:
        json.dump(versions, f, ensure_ascii=False, indent=2)
    log("MODEL", f"偵測到新的體積檔案，建立 {version['version']}（{version['model_name']}）")
    return versions


def cost_of(elements, prices):
    fallback = prices.get("280") or next(iter(prices.values()))
    return sum(el["volume"] * prices.get(el["grade"], fallback) for el in elements.values())


# ---------- 主程式 ----------

def main():
    date = today()
    log("INFO", f"開始估算（{date}）")

    filename, digest, elements = load_quantities()
    versions = update_versions(date, filename, digest, elements)

    prices, source = fetch_market_prices()
    history = update_price_history(date, prices, source)

    # 每個版本都用「目前單價」計算，版本之間才能公平比較
    for v in versions:
        v["total_volume"] = round(sum(e["volume"] for e in v["elements"].values()), 2)
        v["element_count"] = sum(e["count"] for e in v["elements"].values())
        v["cost_at_current_price"] = round(cost_of(v["elements"], prices))

    # 每日成本：當天有效的模型版本 × 當天單價
    cost_history = []
    for d in sorted(history):
        active = [v for v in versions if v["detected_at"] <= d] or versions[:1]
        cost_history.append({"date": d, "version": active[-1]["version"], "cost": round(cost_of(active[-1]["elements"], history[d]))})

    dates = sorted(history)
    previous_prices = history[dates[-2]] if len(dates) > 1 else prices

    latest = versions[-1]
    log("SUCCESS", f"{latest['version']} 總體積 {latest['total_volume']:,.2f} m³，總成本 NT$ {latest['cost_at_current_price']:,}")

    info = {}
    info_path = os.path.join(DATA_DIR, "model_info.json")
    if os.path.exists(info_path):
        with open(info_path, encoding="utf-8") as f:
            info = json.load(f)

    output = {
        "generated_at": datetime.now(TW).strftime("%Y-%m-%d %H:%M"),
        "is_sample": bool(info.get("is_sample")),
        "price_date": date,
        "price_source": source,
        "prices": prices,
        "previous_prices": previous_prices,
        "versions": versions,
        "cost_history": cost_history,
        "logs": logs,
    }
    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2)


if __name__ == "__main__":
    main()
